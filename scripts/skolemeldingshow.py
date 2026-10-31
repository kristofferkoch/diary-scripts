#!/usr/bin/env python3
"""skolemeldingshow — read the Skolemelding JSONL sink as readable text.

The sjekk-track reader for ``skolemelding-sync`` data, in the
``vigiloshown`` / ``spondshow`` shape: header line per record,
``--since-cursor`` / ``--headers-only`` for the collect bundle, ``--full``
when triaging.

Cursor: ``memory/skolemelding-state.json``'s ``cursor`` — **hand-advanced**,
like the mail, Signal and Vigilo cursors, not the auto-bumped
``last_successful_run`` that ``skolemelding-sync`` writes. (``spondshow
--since-cursor`` reads that auto-bumped one, which the sync stamps *after*
fetching — so records from the run that just ran fall just behind it and
vanish from the bundle. Reading a poll cursor as a triage cursor is the bug;
this is why the sync owns ``last_successful_run`` and the human owns
``cursor``.)

Read the SOP in ``scripts/docs/skolemelding.md``: collect → triage →
``skolemeldingshow --since-cursor --bump``.

Examples::

    uv run skolemeldingshow --since-cursor --headers-only   # what the bundle shows
    uv run skolemeldingshow --since-cursor                  # bodies + attachment paths
    uv run skolemeldingshow --since-cursor --full           # + extracted attachment text
    uv run skolemeldingshow --id 12345678 --full            # one message, whole
    uv run skolemeldingshow --all                           # weekly sweep, no cursor
    uv run skolemeldingshow --since-cursor --bump           # triaged, advance cursor
"""
from __future__ import annotations

import argparse
import io
import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterator

from mail_reader.config import workspace_root

PROJECT_ROOT = workspace_root()
STATE_PATH = PROJECT_ROOT / "memory" / "skolemelding-state.json"
JSONL_DIR = PROJECT_ROOT / "memory" / "skolemelding"
ATTACH_DIR = JSONL_DIR / "attachments"

BODY_CHARS = 600
ATTACH_CHARS = 40000
KINDS = ("message", "message_update")


def parse_iso(s: str) -> datetime:
    """Parse ISO-8601 with trailing Z, normalising to UTC.

    >>> parse_iso("2026-10-13T08:00:00Z").isoformat()
    '2026-10-13T08:00:00+00:00'
    >>> parse_iso("2026-10-13").isoformat()
    '2026-10-13T00:00:00+00:00'
    """
    if s.endswith("Z"):
        s = s[:-1] + "+00:00"
    dt = datetime.fromisoformat(s)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def load_state() -> dict[str, Any]:
    if not STATE_PATH.exists():
        return {}
    try:
        data = json.loads(STATE_PATH.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return data if isinstance(data, dict) else {}


def save_state(state: dict[str, Any]) -> None:
    STATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    STATE_PATH.write_text(
        json.dumps(state, indent=2, sort_keys=True, ensure_ascii=False) + "\n"
    )


def iter_records(since: datetime | None, until: datetime | None, kinds: set[str] | None,
                 message_id: str | None) -> Iterator[dict[str, Any]]:
    """Stream JSONL records, oldest file first (``sorted`` on the date names)."""
    if not JSONL_DIR.exists():
        return
    for fp in sorted(JSONL_DIR.glob("*.jsonl")):
        if since is not None:
            try:
                if datetime.strptime(fp.stem, "%Y-%m-%d").replace(
                    tzinfo=timezone.utc
                ).date() < since.date():
                    continue
            except ValueError:
                pass
        for line in fp.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except json.JSONDecodeError:
                continue
            if not isinstance(rec, dict):
                continue
            if kinds and rec.get("kind") not in kinds:
                continue
            when = rec.get("fetched_at")
            if when:
                try:
                    ts = parse_iso(str(when))
                except ValueError:
                    ts = None
                if ts is not None:
                    if since is not None and ts <= since:
                        continue
                    if until is not None and ts > until:
                        continue
            data = rec.get("data") or {}
            if message_id and str(data.get("id")) != message_id:
                continue
            yield rec


# --- Attachment text ---------------------------------------------------------
# Mirrored from vigiloshown (PDF via pypdf, docx via python-docx) — Skolemelding
# attachments are always local files downloaded by the sync, so unlike
# vigiloshown this reader never needs the API at show time.


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join(
        text for page in reader.pages if (text := (page.extract_text() or "").strip())
    )


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(data: bytes, mime: str, name: str) -> str | None:
    """Readable text from an attachment, or None when it is not a text format.

    PDF via pypdf, docx via python-docx (paragraphs then tables, cells joined
    with ``" | "`` so a timetable still reads as a grid), anything ``text/*``
    decoded leniently. Anything else is a file to open by hand, and saying so
    beats pretending.
    """
    mime = (mime or "").lower()
    n = (name or "").lower()
    if mime == "application/pdf" or n.endswith(".pdf"):
        return _pdf_text(data)
    if "wordprocessingml" in mime or "msword" in mime or n.endswith(".docx"):
        return _docx_text(data)
    if mime.startswith("text/") or n.endswith((".txt", ".md", ".csv")):
        return data.decode("utf-8", errors="replace")
    return None


def render_attachment_text(data: bytes, mime: str, name: str) -> str:
    try:
        text = extract_text(data, mime, name)
    except Exception as e:  # noqa: BLE001 — a corrupt PDF/docx is a note, not a crash
        label = name or "vedlegget"
        return f"    note: {label} lot seg ikke tolkes ({type(e).__name__}: {e})"
    if text is None:
        return "    note: ikke-tekstlig (åpne fila manuelt)"
    text = text.strip()
    if len(text) > ATTACH_CHARS:
        text = text[:ATTACH_CHARS] + f"\n…[truncated at {ATTACH_CHARS} chars]"
    return "\n".join(f"    {line}" for line in text.splitlines()) or "    (tomt)"


def local_attachment_path(att: dict[str, Any]) -> Path | None:
    """The file ``skolemelding-sync`` recorded for this attachment, if it is there.

    The sync writes the path it actually used, so the reader never has to
    re-derive a sanitised file name and get it subtly wrong.
    """
    raw = att.get("path")
    if not raw:
        return None
    from scripts.attachment_store import existing_attachment_path
    try:
        return existing_attachment_path(ATTACH_DIR, str(raw))
    except (OSError, ValueError):
        return None


def local_attachment_bytes(att: dict[str, Any]) -> bytes:
    from scripts.attachment_store import read_attachment_path
    return read_attachment_path(ATTACH_DIR, str(att.get("path") or ""))


# --- Rendering ---------------------------------------------------------------


def _trim(s: str, n: int) -> str:
    s = (s or "").replace("\r", " ").replace("\n", " ")
    return s if len(s) <= n else s[: n - 1] + "…"


def render_header(rec: dict[str, Any]) -> str:
    kind = str(rec.get("kind", "?"))
    data = rec.get("data") or {}
    when = str(data.get("timestamp") or rec.get("fetched_at", ""))
    if kind in ("message", "message_update"):
        sender = str(data.get("sender") or "?")
        atts = len(data.get("attachments") or [])
        mark = "U" if kind == "message" else "↩"
        att_note = f" 📎{atts}" if atts else ""
        question = " ❓" if data.get("feedback_question") else ""
        return (
            f"{when} [{mark} {kind:<14}] {sender:<24} "
            f"{_trim(str(data.get('header') or ''), 60)!r}{att_note}{question}"
        )
    return f"{when} [{kind}] {_trim(json.dumps(data, ensure_ascii=False), 120)}"


def render_full(rec: dict[str, Any], *, with_attachments: bool) -> list[str]:
    """One record as text: header, then body, threads, attachment content.

    Without ``--full`` a body is cut at :data:`BODY_CHARS` — enough to triage
    a whole bundle without drowning in ukeplaner; the file path is printed so
    the full thing is one command away.
    """
    out = [render_header(rec)]
    data = rec.get("data") or {}
    limit = 0 if with_attachments else BODY_CHARS

    body = str(data.get("body") or "")
    if limit and len(body) > limit:
        body = body[:limit] + f" …[+{len(body) - limit} chars — --full]"
    out.extend(f"  {line}" for line in body.splitlines())

    for thread in data.get("threads") or []:
        if not isinstance(thread, dict):
            continue
        out.append(f"  --- {thread.get('sender') or '?'} "
                   f"({thread.get('timestamp') or ''}) ---")
        out.extend(f"  {line}" for line in str(thread.get("body") or "").splitlines())
        for name in thread.get("attachments") or []:
            out.append(f"  📎 {name} (på trådsvar — åpne i portalen)")

    for att in data.get("attachments") or []:
        if not isinstance(att, dict):
            continue
        out.append(f"  📎 {att.get('name')} ({att.get('mimeType') or 'ukjent'})")
        path = local_attachment_path(att)
        if path is None:
            note = "    note: ikke lastet ned (prøv ny skolemelding-sync)"
            if att.get("source"):
                note += f" — {att['source']}"
            out.append(note)
            continue
        out.append(f"    file: {path}")
        if not with_attachments:
            continue
        try:
            out.append(render_attachment_text(
                local_attachment_bytes(att), str(att.get("mimeType") or ""),
                str(att.get("name") or "")))
        except (OSError, ValueError) as e:
            out.append(f"    !! kunne ikke lese fila: {e}")
    if data.get("note"):
        out.append(f"  !! {data['note']}")
    return out


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--since-cursor", action="store_true",
                    help="Only records fetched after memory/skolemelding-state.json's cursor.")
    ap.add_argument("--since", default=None, help="ISO date/datetime (strictly after).")
    ap.add_argument("--until", default=None,
                    help="ISO date/datetime (at or before) — pair with --since for a window.")
    ap.add_argument("--kind", action="append", choices=list(KINDS), default=None,
                    help="Restrict to a kind (repeatable).")
    ap.add_argument("--id", default=None, metavar="ID",
                    help="Only the message with this Id (detail view).")
    ap.add_argument("--all", action="store_true",
                    help="Ignore the cursor entirely — the weekly sweep.")
    ap.add_argument("--headers-only", action="store_true",
                    help="One line per record; no bodies, no attachments.")
    ap.add_argument("--full", action="store_true",
                    help="Also extract text from PDF/docx/text attachments (local "
                         "files — no network).")
    ap.add_argument("--bump", action="store_true",
                    help="After printing, advance the cursor to the newest record shown.")
    args = ap.parse_args(argv)

    state = load_state()
    if state.get("auth_required"):
        print(
            "!! SKOLEMELDING AUTH REQUIRED — innloggingen er utløpt. Kjør:\n"
            "!!     uv run --project diary-scripts skolemelding-login",
            file=sys.stderr,
        )

    since: datetime | None = None
    if args.all:
        print("(--all: ignoring the cursor)", file=sys.stderr)
    elif args.since_cursor:
        raw = state.get("cursor")
        if raw:
            since = parse_iso(str(raw))
            print(f"(cursor: {since.isoformat()})", file=sys.stderr)
        else:
            # Line shape is load-bearing: sjekk_collect forgives exit 1 only
            # when every output line matches its "(cursor: …)" /
            # "(no matching …)" quiet-track pattern — a first run before any
            # --bump (or before the first login) must read as a quiet track,
            # not a FAILED one.
            print("(cursor: none yet — showing everything; --bump to set it)",
                  file=sys.stderr)
    elif args.since:
        since = parse_iso(args.since)

    until = parse_iso(args.until) if args.until else None
    kinds = set(args.kind) if args.kind else None

    rows: list[dict[str, Any]] = []
    for rec in iter_records(since, until, kinds, args.id):
        if args.headers_only:
            print(render_header(rec))
        else:
            print("\n".join(render_full(rec, with_attachments=args.full)))
            print()
        rows.append(rec)

    if not rows:
        print("(no matching records)", file=sys.stderr)

    if args.bump and rows:
        newest = max(str(r.get("fetched_at") or "") for r in rows)
        if newest:
            state["cursor"] = newest
            save_state(state)
            print(f"(cursor bumped to {newest})", file=sys.stderr)

    return 0 if rows else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))

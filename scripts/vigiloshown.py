#!/usr/bin/env python3
"""vigiloshown — read the Vigilo JSONL sink as readable text.

The sjekk-track reader for ``vigilo-sync`` data, in the ``spondshow`` /
``signalshow`` shape: header line per record, ``--since-cursor`` /
``--headers-only`` for the collect bundle, ``--full`` when triaging.

Cursor: ``memory/vigilo-state.json``'s ``cursor`` — **hand-advanced**, like
the mail and Signal cursors, not the auto-bumped ``last_successful_run`` that
``vigilo-sync`` writes. (``spondshow --since-cursor`` reads that auto-bumped
one, which the sync stamps *after* fetching — so records from the run that
just ran fall just behind it and vanish from the bundle. Reading a poll
cursor as a triage cursor is the bug; this is why Vigilo separates them: the
sync owns ``last_successful_run``, the human owns ``cursor``.)

Read the SOP in ``scripts/README.md`` → "Vigilo": collect → triage →
``vigiloshown --since-cursor --bump``.

Examples::

    uv run vigiloshown --since-cursor --headers-only   # what the bundle shows
    uv run vigiloshown --since-cursor                  # bodies + attachment paths
    uv run vigiloshown --since-cursor --full           # + extracted attachment text
    uv run vigiloshown --kind consent --since 2026-10-01
    uv run vigiloshown --kind check_out_note --all    # hentemeldinger: window + expiry
    uv run vigiloshown --all --kind news               # weekly sweep, no cursor
    uv run vigiloshown --child aks --full
    uv run vigiloshown --since-cursor --bump           # triaged, advance cursor
"""
from __future__ import annotations

import argparse
import io
import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterator

from mail_reader.config import workspace_root

from scripts.vigilo_api import VigiloApi, VigiloError
from scripts.vigilo_sync import CHECKOUT_EXPIRY_WARN_DAYS

PROJECT_ROOT = workspace_root()
STATE_PATH = PROJECT_ROOT / "memory" / "vigilo-state.json"
JSONL_DIR = PROJECT_ROOT / "memory" / "vigilo"
ATTACH_DIR = JSONL_DIR / "attachments"

BODY_CHARS = 600
ATTACH_CHARS = 40000
KINDS = ("message", "message_update", "news", "consent", "check_out_note")


def parse_iso(s: str) -> datetime:
    """Parse ISO-8601 with trailing Z, normalising to UTC.

    >>> parse_iso("2026-10-11T08:00:00Z").isoformat()
    '2026-10-11T08:00:00+00:00'
    >>> parse_iso("2026-10-11").isoformat()
    '2026-10-11T00:00:00+00:00'
    """
    if s.endswith("Z"):
        s = s[:-1] + "+00:00"
    dt = datetime.fromisoformat(s)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def load_state() -> dict[str, Any]:
    if not STATE_PATH.exists():
        return {}
    try:
        data = json.loads(STATE_PATH.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return data if isinstance(data, dict) else {}


def save_state(state: dict[str, Any]) -> None:
    STATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    STATE_PATH.write_text(
        json.dumps(state, indent=2, sort_keys=True, ensure_ascii=False) + "\n"
    )


def iter_records(since: datetime | None, until: datetime | None, kinds: set[str] | None,
                 child: str | None) -> Iterator[dict[str, Any]]:
    """Stream JSONL records, oldest file first (``sorted`` on the date names)."""
    if not JSONL_DIR.exists():
        return
    for fp in sorted(JSONL_DIR.glob("*.jsonl")):
        if since is not None:
            try:
                if datetime.strptime(fp.stem, "%Y-%m-%d").replace(
                    tzinfo=timezone.utc
                ).date() < since.date():
                    continue
            except ValueError:
                pass
        for line in fp.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except json.JSONDecodeError:
                continue
            if not isinstance(rec, dict):
                continue
            if kinds and rec.get("kind") not in kinds:
                continue
            when = rec.get("fetched_at")
            if when:
                try:
                    ts = parse_iso(str(when))
                except ValueError:
                    ts = None
                if ts is not None:
                    if since is not None and ts <= since:
                        continue
                    if until is not None and ts > until:
                        continue
            data = rec.get("data") or {}
            if child:
                hay = " ".join(
                    str(data.get(k) or "") for k in ("childId", "unit", "childLabel")
                ).lower()
                if child.lower() not in hay:
                    continue
            yield rec


# --- Attachment text ---------------------------------------------------------


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join(
        text for page in reader.pages if (text := (page.extract_text() or "").strip())
    )


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(data: bytes, mime: str, name: str) -> str | None:
    """Readable text from an attachment, or None when it is not a text format.

    PDF via pypdf, docx via python-docx (paragraphs then tables, cells joined
    with ``" | "`` so a timetable still reads as a grid), anything ``text/*``
    decoded leniently. Anything else is a file to open by hand, and saying so
    beats pretending.
    """
    mime = (mime or "").lower()
    n = (name or "").lower()
    if mime == "application/pdf" or n.endswith(".pdf"):
        return _pdf_text(data)
    if "wordprocessingml" in mime or "msword" in mime or n.endswith(".docx"):
        return _docx_text(data)
    if mime.startswith("text/") or n.endswith((".txt", ".md", ".csv")):
        return data.decode("utf-8", errors="replace")
    return None


def render_attachment_text(data: bytes, mime: str, name: str) -> str:
    try:
        text = extract_text(data, mime, name)
    except Exception as e:  # noqa: BLE001 — a corrupt PDF/docx is a note, not a crash
        label = name or "vedlegget"
        return f"    note: {label} lot seg ikke tolkes ({type(e).__name__}: {e})"
    if text is None:
        return "    note: ikke-tekstlig (åpne fila manuelt)"
    text = text.strip()
    if len(text) > ATTACH_CHARS:
        text = text[:ATTACH_CHARS] + f"\n…[truncated at {ATTACH_CHARS} chars]"
    return "\n".join(f"    {line}" for line in text.splitlines()) or "    (tomt)"


def local_attachment_path(att: dict[str, Any]) -> Path | None:
    """The file ``vigilo-sync`` recorded for this attachment, if it is there.

    The sync writes the path it actually used, so the reader never has to
    re-derive a sanitised file name and get it subtly wrong.
    """
    raw = att.get("path")
    if not raw:
        return None
    path = Path(str(raw))
    return path if path.exists() else None


# --- Rendering ---------------------------------------------------------------


def _trim(s: str, n: int) -> str:
    s = (s or "").replace("\r", " ").replace("\n", " ")
    return s if len(s) <= n else s[: n - 1] + "…"


def render_header(rec: dict[str, Any]) -> str:
    kind = str(rec.get("kind", "?"))
    data = rec.get("data") or {}
    when = str(data.get("created") or rec.get("fetched_at", ""))
    unit = str(data.get("unit") or "")
    if kind in ("message", "message_update"):
        msgs = data.get("messages") or []
        last = msgs[-1] if msgs else {}
        sender = str(last.get("sender") or "?")
        atts = sum(len(m.get("attachments") or []) for m in msgs if isinstance(m, dict))
        mark = "U" if kind == "message" else "↩"
        att_note = f" 📎{atts}" if atts else ""
        return (
            f"{when} [{mark} {kind:<14}] {unit:<18} {sender:<22} "
            f"{_trim(str(data.get('title') or ''), 60)!r}{att_note}"
        )
    if kind == "news":
        atts = data.get("attachments") or []
        att_note = f" 📎{len(atts)}" if atts else ""
        return (
            f"{when} [· news          ] {unit:<18} "
            f"{_trim(str(data.get('title') or ''), 70)!r}{att_note}"
        )
    if kind == "consent":
        deadline = data.get("deadline") or "-"
        status = data.get("status") or "?"
        return (
            f"{when} [! consent       ] {unit:<18} "
            f"{_trim(str(data.get('title') or ''), 46)!r} status={status} deadline={deadline}"
        )
    if kind == "check_out_note":
        # The header has to answer the two questions a hentemelding raises:
        # what is the window, and how much of it is left.
        bits: list[str] = []
        if data.get("timeOfDay"):
            bits.append(str(data["timeOfDay"])[:5])
        if data.get("days"):
            bits.append(str(data["days"]))
        if data.get("toDate"):
            bits.append(f"til {data['toDate']}")
        left = data.get("daysLeft")
        if isinstance(left, int) and not isinstance(left, bool) \
                and left <= CHECKOUT_EXPIRY_WARN_DAYS:
            bits.insert(0, "⚠")
        tail = " ".join(filter(None, [" ".join(bits), str(data.get("change") or "")]))
        return (
            f"{when} [! {kind:<14}] {unit:<18} "
            f"{_trim(str(data.get('title') or ''), 44)!r} {tail}".rstrip()
        )
    return f"{when} [{kind}] {_trim(json.dumps(data, ensure_ascii=False), 120)}"


def render_full(rec: dict[str, Any], *, with_attachments: bool,
                api: VigiloApi | None) -> list[str]:
    """One record as text: header, then bodies, then attachment content.

    Without ``--full`` a body is cut at :data:`BODY_CHARS` — enough to triage
    a whole bundle without drowning in ukeplaner; the file path is printed so
    the full thing is one command away.
    """
    out = [render_header(rec)]
    data = rec.get("data") or {}
    kind = str(rec.get("kind"))
    limit = 0 if with_attachments else BODY_CHARS

    if kind in ("message", "message_update"):
        for msg in data.get("messages") or []:
            if not isinstance(msg, dict):
                continue
            out.append(f"  --- {msg.get('sender') or '?'} ({msg.get('timeCreated') or ''}) ---")
            body = str(msg.get("body") or "")
            if limit and len(body) > limit:
                body = body[:limit] + f" …[+{len(body) - limit} chars — --full]"
            out.extend(f"  {line}" for line in body.splitlines())
            for att in msg.get("attachments") or []:
                if not isinstance(att, dict):
                    continue
                out.append(f"  📎 {att.get('name')} ({att.get('mimeType')})")
                path = local_attachment_path(att)
                if path is None:
                    out.append("    note: ikke lastet ned (prøv ny vigilo-sync)")
                    continue
                out.append(f"    file: {path}")
                if not with_attachments:
                    continue
                try:
                    out.append(render_attachment_text(
                        path.read_bytes(), str(att.get("mimeType") or ""),
                        str(att.get("name") or "")))
                except OSError as e:
                    out.append(f"    !! kunne ikke lese fila: {e}")
        if data.get("note"):
            out.append(f"  !! {data['note']}")
    elif kind == "news":
        if data.get("body"):
            out.extend(f"  {line}" for line in str(data["body"]).splitlines())
        for att in data.get("attachments") or []:
            out.append(f"  📎 {att.get('name')} ({att.get('mimeType')})")
            if not with_attachments:
                continue
            if api is None:
                out.append("    note: ikke forsøkt (ingen token-fil)")
                continue
            try:
                blob = api.post_attachment(str(data.get("postId")), str(att.get("id")))
                out.append(render_attachment_text(blob, str(att.get("mimeType") or ""),
                                                  str(att.get("name") or "")))
            except (VigiloError, OSError) as e:
                out.append(f"    !! nedlasting feilet: {e}")
    elif kind == "consent":
        if data.get("body"):
            out.extend(f"  {line}" for line in str(data["body"]).splitlines())
        out.append(f"  childId={data.get('childId')} consentId={data.get('consentId')}")
    elif kind == "check_out_note":
        title = str(data.get("title") or "")
        if title:
            out.extend(f"  {line}" for line in title.splitlines())
        # The verdict is the point of the record ("ny", "endret", "utløper om
        # 12 dager") — the header shows it, so --full must not drop it.
        change = str(data.get("change") or "")
        if change:
            out.append(f"  {change}")
        if data.get("toDate") or data.get("timeOfDay"):
            out.append(
                f"  {data.get('timeOfDay') or '?'} [{data.get('days') or '-'}]  "
                f"{data.get('fromDate') or '?'} → {data.get('toDate') or '?'}"
            )
            left = data.get("daysLeft")
            if isinstance(left, int) and not isinstance(left, bool):
                out.append(f"  {left} dager igjen av gyldigheten" if left >= 0
                           else f"  utløpt for {-left} dager siden")
        if any(data.get(k) is not None
               for k in ("typeName", "category", "enabled", "recurrenceType")):
            recurrence = str(data.get("recurrenceType") or "?")
            interval = data.get("recurrenceInterval")
            if isinstance(interval, int) and interval > 1:
                # Only a multiple is worth printing: "weekly ×1" reads like it
                # means something beside "weekly".
                recurrence += f" ×{interval}"
            out.append(f"  type={data.get('typeName')} category={data.get('category')} "
                       f"enabled={data.get('enabled')} "
                       f"recurrence={recurrence} recurring={data.get('recurring')}")
        out.append(f"  childId={data.get('childId')} noteId={data.get('noteId')}")
    return out


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--since-cursor", action="store_true",
                    help="Only records fetched after memory/vigilo-state.json's cursor.")
    ap.add_argument("--since", default=None, help="ISO date/datetime (strictly after).")
    ap.add_argument("--until", default=None,
                    help="ISO date/datetime (at or before) — pair with --since for a window.")
    ap.add_argument("--kind", action="append", choices=list(KINDS), default=None,
                    help=f"Restrict to a kind (repeatable). check_out_note is the "
                         f"hentemelding with its validity end date; ⚠ marks one "
                         f"inside {CHECKOUT_EXPIRY_WARN_DAYS} days of it.")
    ap.add_argument("--child", default=None, metavar="NEEDLE",
                    help="Only records whose childId / unit contains this (case-insensitive).")
    ap.add_argument("--all", action="store_true",
                    help="Ignore the cursor entirely — the weekly sweep.")
    ap.add_argument("--headers-only", action="store_true",
                    help="One line per record; no bodies, no attachments.")
    ap.add_argument("--full", action="store_true",
                    help="Also extract text from PDF/docx/text attachments (message "
                         "attachments from disk, news attachments over the API).")
    ap.add_argument("--bump", action="store_true",
                    help="After printing, advance the cursor to the newest record shown.")
    args = ap.parse_args(argv)

    state = load_state()
    if state.get("auth_required"):
        print(
            "!! VIGILO AUTH REQUIRED — refresh-tokenet er utløpt. Kjør:\n"
            "!!     uv run --project diary-scripts vigilo-login",
            file=sys.stderr,
        )

    since: datetime | None = None
    if args.all:
        print("(--all: ignoring the cursor)", file=sys.stderr)
    elif args.since_cursor:
        raw = state.get("cursor")
        if raw:
            since = parse_iso(str(raw))
            print(f"(cursor: {since.isoformat()})", file=sys.stderr)
        else:
            # Line shape is load-bearing: sjekk_collect forgives exit 1 only
            # when every output line matches its "(cursor: …)" /
            # "(no matching …)" quiet-track pattern — a first run before any
            # --bump (or before the first login) must read as a quiet track,
            # not a FAILED one.
            print("(cursor: none yet — showing everything; --bump to set it)",
                  file=sys.stderr)
    elif args.since:
        since = parse_iso(args.since)

    until = parse_iso(args.until) if args.until else None
    kinds = set(args.kind) if args.kind else None

    # Only opened when attachment text is actually wanted — a plain header
    # listing must work with a dead token, and must not hit the network.
    api: VigiloApi | None = None
    if args.full and not args.headers_only:
        try:
            api = VigiloApi()
        except VigiloError as e:
            print(f"!! vigilo: ingen token, hopper over vedleggstekst ({e})",
                  file=sys.stderr)

    rows: list[dict[str, Any]] = []
    try:
        for rec in iter_records(since, until, kinds, args.child):
            if args.headers_only:
                print(render_header(rec))
            else:
                print("\n".join(render_full(rec, with_attachments=args.full, api=api)))
                print()
            rows.append(rec)
    finally:
        if api is not None:
            api.close()

    if not rows:
        print("(no matching records)", file=sys.stderr)

    if args.bump and rows:
        newest = max(str(r.get("fetched_at") or "") for r in rows)
        if newest:
            state["cursor"] = newest
            save_state(state)
            print(f"(cursor bumped to {newest})", file=sys.stderr)

    return 0 if rows else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
